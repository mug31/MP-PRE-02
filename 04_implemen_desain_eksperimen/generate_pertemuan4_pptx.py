import os
import sys
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# ============================================================
# DESIGN TOKENS — SWISS ACADEMIC / FOREST GREEN STYLE
# ============================================================
DARK_GREEN = RGBColor(0x1E, 0x42, 0x20)  # #1E4220 (Cover bg, Section headers)
MID_GREEN  = RGBColor(0x2C, 0x5F, 0x2D)  # #2C5F2D (Pill badges, accents)
SAGE_GREEN = RGBColor(0x97, 0xBC, 0x62)  # #97BC62 (Accents, rings)
MINT_LIGHT = RGBColor(0xE4, 0xEF, 0xD8)  # #E4EFD8 (Light green bg)
AMBER_GOLD = RGBColor(0xE8, 0xA3, 0x3D)  # #E8A33D (Gold badges & highlights)
CARD_BG    = RGBColor(0xF4, 0xF7, 0xF1)  # #F4F7F1 (Card background)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)  # #FFFFFF
DARK_TEXT  = RGBColor(0x23, 0x33, 0x24)  # #233324 (Body text)
MUTED_TEXT = RGBColor(0x5B, 0x6B, 0x5C)  # #5B6B5C (Footers & captions)
BORDER_COL = RGBColor(0xD0, 0xDE, 0xCC)  # Border color for cards

FONT_TITLE = "Cambria"
FONT_BODY  = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

def add_header(slide, tag_text, title_text, subtitle_text=""):
    # Top Tag
    tb = slide.shapes.add_textbox(Inches(0.80), Inches(0.45), Inches(11.00), Inches(0.35))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = tag_text.upper()
    p.font.name = FONT_BODY
    p.font.size = Pt(10.0)
    p.font.bold = True
    p.font.color.rgb = AMBER_GOLD

    # Title
    tb_t = slide.shapes.add_textbox(Inches(0.80), Inches(0.75), Inches(11.50), Inches(0.65))
    tf_t = tb_t.text_frame
    tf_t.word_wrap = True
    tf_t.margin_left = tf_t.margin_top = tf_t.margin_right = tf_t.margin_bottom = 0
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.name = FONT_TITLE
    p_t.font.size = Pt(22.0)
    p_t.font.bold = True
    p_t.font.color.rgb = DARK_GREEN

    if subtitle_text:
        tb_s = slide.shapes.add_textbox(Inches(0.80), Inches(1.35), Inches(11.50), Inches(0.35))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = tf_s.margin_top = tf_s.margin_right = tf_s.margin_bottom = 0
        p_s = tf_s.paragraphs[0]
        p_s.text = subtitle_text
        p_s.font.name = FONT_BODY
        p_s.font.size = Pt(11.5)
        p_s.font.color.rgb = MUTED_TEXT

def add_footer(slide, page_num):
    tx_box = slide.shapes.add_textbox(Inches(0.80), Inches(7.10), Inches(8.00), Inches(0.30))
    tf = tx_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "PRE-02 — Prediksi Probabilitas Keterlambatan Proyek Multi Sumber Daya Terbatas"
    p.font.name = FONT_BODY
    p.font.size = Pt(9.0)
    p.font.color.rgb = MUTED_TEXT

    tx_num = slide.shapes.add_textbox(Inches(12.00), Inches(7.10), Inches(0.50), Inches(0.30))
    tf_num = tx_num.text_frame
    tf_num.word_wrap = True
    tf_num.margin_left = tf_num.margin_top = tf_num.margin_right = tf_num.margin_bottom = 0
    p_num = tf_num.paragraphs[0]
    p_num.text = str(page_num)
    p_num.font.name = FONT_BODY
    p_num.font.size = Pt(9.0)
    p_num.font.color.rgb = MUTED_TEXT

# ============================================================
# SLIDE 1: COVER
# ============================================================
s1 = prs.slides.add_slide(blank_layout)
s1.background.fill.solid()
s1.background.fill.fore_color.rgb = DARK_GREEN

# Decorative circles
e1 = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.50), Inches(4.50), Inches(5.00), Inches(5.00))
e1.fill.solid()
e1.fill.fore_color.rgb = MID_GREEN
e1.line.fill.background()

e2 = s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(11.30), Inches(5.30), Inches(3.40), Inches(3.40))
e2.fill.solid()
e2.fill.fore_color.rgb = SAGE_GREEN
e2.line.fill.background()

tb = s1.shapes.add_textbox(Inches(0.80), Inches(1.30), Inches(8.00), Inches(0.40))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = "PROGRESS PENELITIAN PERTEMUAN 4"
p.font.name = FONT_BODY
p.font.size = Pt(14.0)
p.font.bold = True
p.font.color.rgb = AMBER_GOLD

tb = s1.shapes.add_textbox(Inches(0.80), Inches(1.80), Inches(10.50), Inches(2.20))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Implementasi Desain Eksperimen:\nPemodelan Probabilitas Terkalibrasi & Early Warning System"
p.font.name = FONT_TITLE
p.font.size = Pt(32.0)
p.font.bold = True
p.font.color.rgb = WHITE

tb = s1.shapes.add_textbox(Inches(0.80), Inches(4.20), Inches(10.00), Inches(1.00))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Topik PRE-02: Prediksi Probabilitas Keterlambatan Proyek Multi Sumber Daya Terbatas\nStudi Kasus: Portofolio Modul Pengembangan Super ERP Farm Nation Enterprise 2026"
p.font.name = FONT_BODY
p.font.size = Pt(14.0)
p.font.color.rgb = MINT_LIGHT

tb = s1.shapes.add_textbox(Inches(0.80), Inches(5.80), Inches(8.00), Inches(0.80))
tf = tb.text_frame
p = tf.paragraphs[0]
p.text = "Tim Peneliti PRE-02 | Mata Kuliah Manajemen Proyek 2026"
p.font.name = FONT_BODY
p.font.size = Pt(12.0)
p.font.color.rgb = SAGE_GREEN

# ============================================================
# SLIDE 2: REKAP KESIAPAN DATA MENUJU EKSPERIMEN
# ============================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Kesiapan Fondasi Data (Pertemuan 3 Recap)", "Pondasi Empiris: Dataset Portofolio Modul Super ERP FNE", "26 aktivitas inti modul ERP diekstraksi dari 8 dokumen PMP untuk pengujian eksperimen.")
add_footer(s2, 2)

cards_s2 = [
    ("Sumber Data PMP", "Data diekstraksi dari 8 sub-proyek FNE (P-FNE-01 Foundation s.d P-FNE-08 Autonomous Enterprise) berstandar PMBOK dan INKINDO 2026.", "8 Dokumen PMP"),
    ("Ukuran Sampel & Distribusi", "N = 26 Aktivitas Operasional\n• Status Terlambat: 16 Modul (61.5%)\n• Status Tepat Waktu: 10 Modul (38.5%)\nDistribusi seimbang secara wajar.", "N = 26 Tasks"),
    ("Fitur Multi-Dimensi", "7 Variabel Masukan (X1..X7):\n• Durasi & Person-Hours\n• Predecessor Count (Dependensi)\n• Resource Utilization Rate\n• Risk Score & SPI Value\n• Scope Change Requests", "7 Prediktor Kunci")
]

for i, (title, body, tag) in enumerate(cards_s2):
    left = Inches(0.80 + i * 4.00)
    top = Inches(2.00)
    width = Inches(3.70)
    height = Inches(4.70)
    
    card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COL
    
    # Tag
    tb = s2.shapes.add_textbox(left + Inches(0.30), top + Inches(0.30), Inches(3.00), Inches(0.35))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = tag.upper()
    p.font.name = FONT_BODY
    p.font.size = Pt(10.0)
    p.font.bold = True
    p.font.color.rgb = MID_GREEN
    
    # Title
    tb = s2.shapes.add_textbox(left + Inches(0.30), top + Inches(0.70), Inches(3.10), Inches(0.60))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_TITLE
    p.font.size = Pt(16.0)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    
    # Body
    tb = s2.shapes.add_textbox(left + Inches(0.30), top + Inches(1.40), Inches(3.10), Inches(3.00))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    p.font.name = FONT_BODY
    p.font.size = Pt(12.0)
    p.font.color.rgb = DARK_TEXT

# ============================================================
# SLIDE 3: FLOW ALUR PIPELINE EKSPERIMEN ML
# ============================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Pipeline Arsitektur", "Flowchart Alur Eksperimen Machine Learning End-to-End", "Rancangan proses standar saintifik adaptasi CRISP-DM untuk predictive project monitoring.")
add_footer(s3, 3)

steps = [
    ("Tahap 1", "Data Ingestion & Audit", "Pembacaan 26 modul FNE & verifikasi integritas 7 fitur input."),
    ("Tahap 2", "Partisi & Penskalaan", "Stratified 5-Fold CV & StandardScaler terisolasi per fold."),
    ("Tahap 3", "Pelatihan 4 Algoritma", "Eksekusi LR Baseline, Random Forest, Gradient Boosting, dan MLP."),
    ("Tahap 4", "Kalibrasi Probabilitas", "Platt Scaling & Isotonic Regression, dilatih pada inner 3-fold di dalam fold training (tanpa kebocoran data)."),
    ("Tahap 5", "Evaluasi Multi-Metrik", "Kalkulasi Brier Score, ROC-AUC, ECE, Log Loss, dan kurva keandalan."),
    ("Tahap 6", "Threshold & EWS", "Analisis ambang Youden J dan translasi probabilitas ke 3 zona mitigasi bagi Project Manager.")
]

for i, (tag, title, desc) in enumerate(steps):
    left = Inches(0.80 + (i % 3) * 4.00)
    top = Inches(2.00 + (i // 3) * 2.45)
    width = Inches(3.70)
    height = Inches(2.20)
    
    card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COL
    
    tb = s3.shapes.add_textbox(left + Inches(0.25), top + Inches(0.20), Inches(3.20), Inches(0.30))
    p = tb.text_frame.paragraphs[0]
    p.text = tag.upper()
    p.font.name = FONT_BODY
    p.font.size = Pt(10.0)
    p.font.bold = True
    p.font.color.rgb = AMBER_GOLD
    
    tb = s3.shapes.add_textbox(left + Inches(0.25), top + Inches(0.50), Inches(3.20), Inches(0.50))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.name = FONT_TITLE
    p.font.size = Pt(14.0)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    
    tb = s3.shapes.add_textbox(left + Inches(0.25), top + Inches(1.05), Inches(3.20), Inches(1.00))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = desc
    p.font.name = FONT_BODY
    p.font.size = Pt(11.0)
    p.font.color.rgb = DARK_TEXT

# ============================================================
# SLIDE 4: 5 SKENARIO EKSPERIMEN ALGORITMA
# ============================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Desain Skenario", "5 Skenario Eksperimen Pemodelan Probabilitas", "Mengombinasikan model parametrik, ensemble non-linier, jaringan saraf tiruan, dan kalibrasi.")
add_footer(s4, 4)

scenarios = [
    ("Skenario 1", "Logistic Regression", "Baseline Model", "Probabilitas via fungsi sigmoid; regularisasi L2 (C = 1.0) diselesaikan hingga konvergen. Berperan sebagai benchmark linear."),
    ("Skenario 2", "Random Forest", "predict_proba", "100 pohon (max_depth 3). Probabilitas = rata-rata proporsi kelas delay di daun seluruh pohon; menghasilkan Feature Importance MDI."),
    ("Skenario 3", "Gradient Boosting", "Loss Optimization", "100 pohon regresi berurutan pada pseudo-residual log-loss (learning rate 0.05, max_depth 2, subsample 0.8)."),
    ("Skenario 4", "MLP Neural Network", "Deep Sigmoid Layer", "2 lapis tersembunyi (16-8 neuron, ReLU) dan output sigmoid; optimasi Adam (lr 0.01, 1000 iterasi, alpha 0.01)."),
    ("Skenario 5", "Kalibrasi & ECE", "Platt / Isotonic", "Tiap model diuji Raw vs Platt vs Isotonic. Kombinasi dengan Brier Score terendah menjadi model rekomendasi.")
]

for i, (tag, name, sub, desc) in enumerate(scenarios):
    left = Inches(0.80 + i * 2.40)
    top = Inches(2.00)
    width = Inches(2.25)
    height = Inches(4.70)
    
    card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COL
    
    tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(0.25), Inches(1.95), Inches(0.30))
    p = tb.text_frame.paragraphs[0]
    p.text = tag.upper()
    p.font.name = FONT_BODY
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = MID_GREEN
    
    tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(0.55), Inches(1.95), Inches(0.65))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = name
    p.font.name = FONT_TITLE
    p.font.size = Pt(14.0)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    
    tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(1.25), Inches(1.95), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = sub
    p.font.name = FONT_BODY
    p.font.size = Pt(10.0)
    p.font.bold = True
    p.font.color.rgb = AMBER_GOLD
    
    tb = s4.shapes.add_textbox(left + Inches(0.15), top + Inches(1.70), Inches(1.95), Inches(2.80))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = desc
    p.font.name = FONT_BODY
    p.font.size = Pt(11.0)
    p.font.color.rgb = DARK_TEXT

# ============================================================
# SLIDE 5: PROTOKOL VALIDASI & PENGENDALIAN OVERFITTING
# ============================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Validasi & Replikasi", "Protokol Kontrol Eksperimen & Penanganan Data Terbatas", "Menjamin integritas ilmiah, mencegah kebocoran data, dan mengunci replikasi hasil.")
add_footer(s5, 5)

cards_s5 = [
    ("Stratified 5-Fold CV", "Mengingat N=26, Stratified K-Fold membagi data menjadi 5 lipatan dengan menjaga rasio 61.5% Delay : 38.5% On-Time. Hold-out 80:20 (21:5 task) dijalankan sebagai pembanding indikatif.", "Strategi Evaluasi"),
    ("Isolasi Standard Scaler", "Penskalaan z-score dan kalibrator dihitung HANYA dari data latih tiap fold (kalibrator via inner 3-fold), lalu diterapkan ke data uji. Mencegah data leakage.", "Pencegahan Leakage"),
    ("Regularisasi & Fixed Seed", "• Penguncian random_state = 42 untuk seluruh model.\n• Pembatasan max_depth: 3 (RF) dan 2 (GB).\n• Penalti L2 pada Logistic Regression & MLP.\n• Implementasi diverifikasi terhadap scikit-learn.", "Replikasi Terjamin")
]

for i, (title, body, tag) in enumerate(cards_s5):
    left = Inches(0.80 + i * 4.00)
    top = Inches(2.00)
    width = Inches(3.70)
    height = Inches(4.70)
    
    card = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COL
    
    tb = s5.shapes.add_textbox(left + Inches(0.30), top + Inches(0.30), Inches(3.00), Inches(0.35))
    p = tb.text_frame.paragraphs[0]
    p.text = tag.upper()
    p.font.name = FONT_BODY
    p.font.size = Pt(10.0)
    p.font.bold = True
    p.font.color.rgb = AMBER_GOLD
    
    tb = s5.shapes.add_textbox(left + Inches(0.30), top + Inches(0.70), Inches(3.10), Inches(0.60))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_TITLE
    p.font.size = Pt(16.0)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    
    tb = s5.shapes.add_textbox(left + Inches(0.30), top + Inches(1.40), Inches(3.10), Inches(3.00))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = body
    p.font.name = FONT_BODY
    p.font.size = Pt(12.0)
    p.font.color.rgb = DARK_TEXT

# ============================================================
# SLIDE 6: METRIK EVALUASI DUA DIMENSI
# ============================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Kriteria Evaluasi", "Metrik Pengujian Kinerja Probabilistik & Diskriminasi", "Menilai tidak hanya ketepatan klasifikasi, tetapi juga tingkat keandalan numerik probabilitas.")
add_footer(s6, 6)

# Left Box: Kualitas Probabilitas
card_left = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(2.00), Inches(5.60), Inches(4.70))
card_left.fill.solid()
card_left.fill.fore_color.rgb = CARD_BG
card_left.line.color.rgb = BORDER_COL

tb = s6.shapes.add_textbox(Inches(1.10), Inches(2.30), Inches(5.00), Inches(0.40))
p = tb.text_frame.paragraphs[0]
p.text = "DIMENSI 1: KUALITAS PROBABILITAS"
p.font.name = FONT_TITLE
p.font.size = Pt(15.0)
p.font.bold = True
p.font.color.rgb = DARK_GREEN

tb = s6.shapes.add_textbox(Inches(1.10), Inches(2.80), Inches(5.00), Inches(3.60))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "1. Brier Score (BS):\nMengukur rata-rata kuadrat deviasi antara probabilitas estimasi dan luaran aktual (0-1). Semakin mendekati 0, semakin presisi.\n\n2. Expected Calibration Error (ECE):\nMengukur kesenjangan absolut antara confidence level prediksi dan empirical accuracy dalam 5 bin interval (N = 26).\n\n3. Logarithmic Loss (Cross-Entropy):\nMemberikan penalti tajam terhadap prediksi probabilitas yang terlalu percaya diri (overconfident) namun salah."
p.font.name = FONT_BODY
p.font.size = Pt(11.5)
p.font.color.rgb = DARK_TEXT

# Right Box: Kemampuan Diskriminasi
card_right = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.80), Inches(2.00), Inches(5.70), Inches(4.70))
card_right.fill.solid()
card_right.fill.fore_color.rgb = CARD_BG
card_right.line.color.rgb = BORDER_COL

tb = s6.shapes.add_textbox(Inches(7.10), Inches(2.30), Inches(5.10), Inches(0.40))
p = tb.text_frame.paragraphs[0]
p.text = "DIMENSI 2: KEMAMPUAN DISKRIMINASI"
p.font.name = FONT_TITLE
p.font.size = Pt(15.0)
p.font.bold = True
p.font.color.rgb = MID_GREEN

tb = s6.shapes.add_textbox(Inches(7.10), Inches(2.80), Inches(5.10), Inches(3.60))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "1. ROC-AUC (Area Under ROC Curve):\nMengukur kapasitas pemisahan modul terlambat vs tepat waktu di semua kemungkinan nilai threshold.\n\n2. Sensitivity (Recall) & Specificity:\nMenilai kemampuan sistem mendeteksi keterlambatan nyata sekaligus meminimalkan alarm palsu (false alarm).\n\n3. F1-Score & Accuracy:\nEvaluasi performa klasifikasi pada threshold 0.5, dilengkapi ambang optimal Youden J."
p.font.name = FONT_BODY
p.font.size = Pt(11.5)
p.font.color.rgb = DARK_TEXT

# ============================================================
# SLIDE 7: EARLY WARNING SYSTEM ZONASI KEPUTUSAN
# ============================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Aplikasi Manajerial", "Early Warning System (EWS): Translasi Probabilitas ke Mitigasi", "Menerjemahkan taksiran probabilitas numerik menjadi kebijakan operasional nyata bagi Project Manager.")
add_footer(s7, 7)

zones = [
    ("Zona Hijau: Normal / Low Risk", "P(Delay) < 0.35", "Kondisi stabil terkendali. Tidak diperlukan alokasi developer tambahan. Monitoring rutin mingguan.", "#2C5F2D"),
    ("Zona Kuning: Moderate Risk / Watchlist", "0.35 <= P(Delay) < 0.65", "Modul rawan terdampak cascading delay. Kunci scope changes, audit dependensi predecessor, koordinasi harian.", "#E8A33D"),
    ("Zona Merah: Critical Alert / High Risk", "P(Delay) >= 0.65", "Risiko keterlambatan sangat pasti. Tindakan korektif agresif: crashing alokasi developer, fast-tracking rilis, eskalasi ke Steering Committee.", "#B33927")
]

for i, (title, rng, action, color_hex) in enumerate(zones):
    top = Inches(2.00 + i * 1.65)
    left = Inches(0.80)
    width = Inches(11.70)
    height = Inches(1.45)
    
    card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD_BG
    card.line.color.rgb = BORDER_COL
    
    # Left accent bar
    bar = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.20), height)
    bar.fill.solid()
    if i == 0:
        bar.fill.fore_color.rgb = MID_GREEN
    elif i == 1:
        bar.fill.fore_color.rgb = AMBER_GOLD
    else:
        bar.fill.fore_color.rgb = RGBColor(0xB3, 0x39, 0x27)
    bar.line.fill.background()
    
    # Text
    tb = s7.shapes.add_textbox(left + Inches(0.45), top + Inches(0.15), Inches(7.50), Inches(0.40))
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.name = FONT_TITLE
    p.font.size = Pt(15.0)
    p.font.bold = True
    p.font.color.rgb = DARK_GREEN
    
    # Range badge
    tb_r = s7.shapes.add_textbox(left + Inches(8.50), top + Inches(0.15), Inches(2.90), Inches(0.40))
    p_r = tb_r.text_frame.paragraphs[0]
    p_r.text = rng
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.font.name = FONT_BODY
    p_r.font.size = Pt(13.0)
    p_r.font.bold = True
    p_r.font.color.rgb = bar.fill.fore_color.rgb
    
    # Action text
    tb_a = s7.shapes.add_textbox(left + Inches(0.45), top + Inches(0.60), Inches(10.80), Inches(0.70))
    tf_a = tb_a.text_frame
    tf_a.word_wrap = True
    p_a = tf_a.paragraphs[0]
    p_a.text = f"Kebijakan Manajerial: {action}"
    p_a.font.name = FONT_BODY
    p_a.font.size = Pt(11.5)
    p_a.font.color.rgb = DARK_TEXT

# ============================================================
# DATA HASIL EKSPERIMEN (dibaca dari hasil_eksperimen/, jalankan experiment_pipeline.py terlebih dahulu)
# ============================================================
script_dir = os.path.dirname(os.path.abspath(__file__))
result_dir = os.path.join(script_dir, "hasil_eksperimen")
sys.path.insert(0, script_dir)
import experiment_pipeline as ep

def read_csv(name):
    with open(os.path.join(result_dir, name), encoding="utf-8") as f:
        return list(csv.DictReader(f))

metric_rows = read_csv("tabel_metrik_evaluasi.csv")
for r in metric_rows:
    for k in ("ROC_AUC", "Brier_Score", "Log_Loss", "ECE", "Accuracy"):
        r[k] = float(r[k])
best = min(metric_rows, key=lambda r: (round(r["Brier_Score"], 6), round(r["ECE"], 6), r["Log_Loss"]))
best_label = best["Model"] + (" + " + best["Kalibrasi"] if best["Kalibrasi"] != "Raw" else " (tanpa kalibrasi)")

pred_rows = read_csv("tabel_prediksi_probabilitas_task.csv")
y_true = [int(r["Status_Delay_Aktual"]) for r in pred_rows]
p_best = [float(r["P_Delay_Rekomendasi"]) for r in pred_rows]
theta, j_stat, sens, spec = ep.youden_threshold(y_true, p_best)
n_delay = sum(y_true)
red_hit = sum(1 for y, p in zip(y_true, p_best) if p >= ep.EWS_HIGH and y == 1)
red_false = sum(1 for y, p in zip(y_true, p_best) if p >= ep.EWS_HIGH and y == 0)
green_miss = sum(1 for y, p in zip(y_true, p_best) if p < ep.EWS_LOW and y == 1)

fi_rows = read_csv("tabel_feature_importance.csv")
raw_rows = [r for r in metric_rows if r["Kalibrasi"] == "Raw"]
max_auc = max(r["ROC_AUC"] for r in raw_rows)
top_auc = [r["Model"] for r in raw_rows if abs(r["ROC_AUC"] - max_auc) < 1e-9]

def brier_of(model, method):
    return next(r["Brier_Score"] for r in metric_rows if r["Model"] == model and r["Kalibrasi"] == method)

calib_helped = [r["Model"] for r in raw_rows
                if min(brier_of(r["Model"], "Platt"), brier_of(r["Model"], "Isotonic")) < r["Brier_Score"]]

def set_cell(cell, text, size=10.5, bold=False, color=DARK_TEXT, fill=None, align=PP_ALIGN.CENTER):
    cell.text = text
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    p.font.name = FONT_BODY
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    if fill is not None:
        cell.fill.solid()
        cell.fill.fore_color.rgb = fill

# ============================================================
# SLIDE 8: HASIL EVALUASI 5-FOLD CV
# ============================================================
s8r = prs.slides.add_slide(blank_layout)
add_header(s8r, "Hasil Eksperimen", "Kinerja 4 Algoritma × 3 Metode Kalibrasi (Stratified 5-Fold CV)",
           f"Baris tebal = model rekomendasi (Brier Score terendah): {best_label}.")
add_footer(s8r, 8)

headers = ["Model", "Kalibrasi", "ROC-AUC", "Brier Score", "Log Loss", "ECE", "Akurasi"]
col_widths = [3.70, 1.40, 1.30, 1.40, 1.30, 1.20, 1.40]
table = s8r.shapes.add_table(len(metric_rows) + 1, len(headers), Inches(0.80), Inches(1.95),
                             Inches(sum(col_widths)), Inches(0.33 * (len(metric_rows) + 1))).table
for c, w in enumerate(col_widths):
    table.columns[c].width = Inches(w)
for c, h in enumerate(headers):
    set_cell(table.cell(0, c), h, bold=True, color=WHITE, fill=DARK_GREEN)
for i, r in enumerate(metric_rows, start=1):
    is_best = r is best
    fill = MINT_LIGHT if is_best else (CARD_BG if i % 2 else WHITE)
    values = [r["Model"], r["Kalibrasi"], f"{r['ROC_AUC']:.3f}", f"{r['Brier_Score']:.4f}",
              f"{r['Log_Loss']:.3f}", f"{r['ECE']:.3f}", f"{r['Accuracy'] * 100:.1f}%"]
    for c, v in enumerate(values):
        set_cell(table.cell(i, c), v, bold=is_best, fill=fill, align=PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER)

tb = s8r.shapes.add_textbox(Inches(0.80), Inches(6.35), Inches(11.70), Inches(0.60))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = ("Akurasi dihitung pada threshold 0.5. Log Loss Isotonic tinggi karena kalibrator menghasilkan probabilitas tepat 0/1 "
          "pada sebagian task; ini wajar untuk kalibrator non-parametrik yang dilatih pada ±20 sampel per fold.")
p.font.name = FONT_BODY
p.font.size = Pt(10.0)
p.font.color.rgb = MUTED_TEXT

# ============================================================
# SLIDE 9: TEMUAN KUNCI
# ============================================================
s9r = prs.slides.add_slide(blank_layout)
add_header(s9r, "Temuan Kunci", "Apa yang Ditunjukkan Hasil Eksperimen?",
           "Seluruh angka dibaca langsung dari keluaran experiment_pipeline.py.")
add_footer(s9r, 9)

card = s9r.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), Inches(1.95), Inches(6.30), Inches(3.55))
card.fill.solid()
card.fill.fore_color.rgb = CARD_BG
card.line.color.rgb = BORDER_COL

findings = [
    ("Diskriminasi", f"ROC-AUC tertinggi {max_auc:.3f} dicapai {ep.fmt_models(top_auc)}."),
    ("Model rekomendasi", f"{best_label}: Brier {best['Brier_Score']:.4f}, ECE {best['ECE']:.3f}."),
    ("Efek kalibrasi", ("Kalibrasi menurunkan Brier Score pada " + ep.fmt_models(calib_helped) + "; model lain tidak membaik.")
                       if calib_helped else "Kalibrasi tidak menurunkan Brier Score pada model mana pun; output mentah sudah terkalibrasi baik."),
    ("Fitur dominan", f"{fi_rows[0]['Feature']} ({float(fi_rows[0]['Importance_MDI']):.3f}) dan "
                      f"{fi_rows[1]['Feature']} ({float(fi_rows[1]['Importance_MDI']):.3f}) menurut MDI Random Forest."),
    ("Youden J", f"θ* = {theta:.3f} (Sensitivity {sens:.2f}, Specificity {spec:.2f})."),
    ("EWS", f"Zona Merah menangkap {red_hit}/{n_delay} task terlambat dengan {red_false} alarm palsu; "
            f"{green_miss} task terlambat lolos ke zona Hijau."),
]
tb = s9r.shapes.add_textbox(Inches(1.10), Inches(2.15), Inches(5.80), Inches(3.20))
tf = tb.text_frame
tf.word_wrap = True
for i, (label, text) in enumerate(findings):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(7)
    run = p.add_run()
    run.text = label.upper() + "  "
    run.font.name = FONT_BODY
    run.font.size = Pt(10.0)
    run.font.bold = True
    run.font.color.rgb = MID_GREEN
    run = p.add_run()
    run.text = text
    run.font.name = FONT_BODY
    run.font.size = Pt(12.0)
    run.font.color.rgb = DARK_TEXT

s9r.shapes.add_picture(os.path.join(result_dir, "feature_importance_comparison.png"),
                       Inches(7.35), Inches(2.10), width=Inches(5.20))
tb = s9r.shapes.add_textbox(Inches(7.35), Inches(5.20), Inches(5.20), Inches(1.60))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = ("Catatan: fitur multi sumber daya (Resource_Utilization_Rate, Predecessor_Count) bukan prediktor dominan pada data ini, "
          "sehingga hipotesis kendala multi sumber daya perlu dikaji ulang pada Bab IV.")
p.font.name = FONT_BODY
p.font.size = Pt(10.5)
p.font.italic = True
p.font.color.rgb = MUTED_TEXT

# ============================================================
# SLIDE 10: SUMMARY & NEXT STEP MENUJU PERTEMUAN 5
# ============================================================
s8 = prs.slides.add_slide(blank_layout)
s8.background.fill.solid()
s8.background.fill.fore_color.rgb = DARK_GREEN

tb = s8.shapes.add_textbox(Inches(0.80), Inches(1.20), Inches(8.00), Inches(0.40))
p = tb.text_frame.paragraphs[0]
p.text = "STATUS & ROADMAP KELANJUTAN"
p.font.name = FONT_BODY
p.font.size = Pt(13.0)
p.font.bold = True
p.font.color.rgb = AMBER_GOLD

tb = s8.shapes.add_textbox(Inches(0.80), Inches(1.70), Inches(11.50), Inches(1.20))
tf = tb.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Capaian Pertemuan 4:\nDesain Eksperimen Terimplementasi & Arah Pertemuan 5"
p.font.name = FONT_TITLE
p.font.size = Pt(26.0)
p.font.bold = True
p.font.color.rgb = WHITE

cards_s8 = [
    ("1. Deliverables", "Flow & 5 skenario eksperimen, experiment_pipeline.py (4 model + kalibrasi Platt/Isotonic, diverifikasi terhadap scikit-learn), dan draft Bab III Metodologi."),
    ("2. Keterbatasan", "N = 26 dan data empiris-simulatif: ROC-AUC ≈ 1.0 berarti kelas hampir terpisah sempurna, sehingga hasil belum dapat digeneralisasi."),
    ("3. Pertemuan 5", "Validasi pada data proyek riil yang lebih besar, kaji ulang hipotesis multi sumber daya terhadap fitur dominan, dan susun Bab IV.")
]

for i, (title, desc) in enumerate(cards_s8):
    left = Inches(0.80 + i * 4.00)
    top = Inches(3.30)
    width = Inches(3.70)
    height = Inches(3.20)
    
    card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = MID_GREEN
    card.line.fill.background()
    
    tb = s8.shapes.add_textbox(left + Inches(0.30), top + Inches(0.40), Inches(3.10), Inches(0.60))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = FONT_TITLE
    p.font.size = Pt(15.0)
    p.font.bold = True
    p.font.color.rgb = AMBER_GOLD
    
    tb = s8.shapes.add_textbox(left + Inches(0.30), top + Inches(1.10), Inches(3.10), Inches(1.80))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = desc
    p.font.name = FONT_BODY
    p.font.size = Pt(12.0)
    p.font.color.rgb = MINT_LIGHT

output_pptx = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PRE02_Pertemuan4_Implementasi_Desain_Eksperimen.pptx")
prs.save(output_pptx)
print(f"[SUCCESS] Slide Pertemuan 4 berhasil dibuat: {output_pptx}")
